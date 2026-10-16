"""Build the five-page submission report from verified experiment artifacts."""
import csv
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'CA6124 Final Report.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.bold = False
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.04
for name, size in [('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.color.rgb = RGBColor(55,55,55)
doc.core_properties.title = 'Extending DaisyRec for Item Ranking'
doc.core_properties.subject = 'CA6124 recommendation system project'
doc.core_properties.author = '[Student name]'
# Remove decorative paragraph borders inherited from the runtime template.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

def p(text='', style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 1)

def table(headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for i, val in enumerate(headers):
        t.rows[0].cells[i].text = val
    for row in rows:
        for cell, val in zip(t.add_row().cells, row):
            cell.text = str(val)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[ci])
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.space_before = Pt(4)
                for run in para.runs:
                    run.font.size = Pt(size)
                    run.bold = ri == 0
            if ri == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E9EEF3')
                cell._tc.get_or_add_tcPr().append(shading)
    p('').paragraph_format.space_after = Pt(0)
    return t

def page():
    doc.add_page_break()

def code(text):
    para = p(text)
    for run in para.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(9)

p('Extending DaisyRec for Item Ranking', 'Title')
p('CA6124 Recommendation System Project')
p('Name: [Student name]    Student ID: [Student ID]')
h('1 Introduction')
p('This project extends DaisyRec [1] with SimpleX and compares it with Matrix Factorization (MF) and Neural Matrix Factorization (NeuMF) on MovieLens and Yelp subsets. The extension adds training-history aggregation and cosine contrastive learning, while retaining a common data split and ranking protocol. Across six completed final experiments, NeuMF achieved the highest test NDCG@10 on ml-50k (0.120102), while SimpleX achieved the highest on yelp-200k (0.082864). These outcomes apply to the sampled, single-seed protocol used here.')
p('Project repository: [Private GitHub repository URL]')
h('2 Datasets and preparation')
p('MovieLens ml-50k is a sample of 50,000 ratings from GroupLens MovieLens 1M [2], drawn with pandas sample(n=50000, random_state=42). No count filter was applied. A fresh source download reproduced the supplied subset exactly, including row order.')
p('The Yelp source is the Kaggle Yelp dataset [3]. The recorded procedure converts review user_id, business_id, stars and date into interaction records, samples 200,000 reviews with seed 42, and retains rows whose user and item each appeared at least five times in that initial sample. This single-pass filter is not an iterative 5-core. The supplied CSV contains 88,692 rows; removing 914 duplicate user-item pairs leaves 87,778 interactions. The exact historical Yelp release was not recorded, so a new download may produce a different subset.')
p('Table 1. Dataset statistics after experiment preprocessing', 'Caption')
table(['Dataset','Users','Items','Interactions','Sparsity'],[
    ['ml-50k','5,620','3,098','50,000','99.712821%'],
    ['yelp-200k','13,736','12,287','87,778','99.947991%']], [1.3,1.05,1.05,1.6,1.75], 10)
p('Sparsity = 100 × [1 − interactions / (users × items)]. All ratings become implicit positives without a star-rating threshold. Duplicate pairs retain the last occurrence in file order before timestamp sorting; no additional count filtering occurs in the experiment loader.')
h('3 Experimental protocol')
p('A global chronological split holds out the latest 20% for test. Tuning reserves the latest 10% of the remaining 80% for validation, giving approximately 72% fitting, 8% validation and 20% test. After selection, final fitting uses the full 80%. Histories and graphs are reconstructed from the current fitting split, with no held-out test interactions.')
table(['Dataset','Tuning fit','Validation','Final fit','Test'],[
    ['ml-50k','36,000','4,000','40,000','10,000'],
    ['yelp-200k','63,201','7,022','70,223','17,555']], [1.35,1.35,1.35,1.35,1.35], 10)

page()
h('4 Model implementation')
p('MF scores users and items through learned embedding dot products. NeuMF combines generalized matrix factorization with a multilayer perceptron. Both baselines train with Bayesian Personalized Ranking (BPR). SimpleX uses a mean behavior aggregator and Cosine Contrastive Loss (CCL), following the SimpleX design [4].')
p('For user u, let H(u) contain only items in the current training history. The implementation computes:')
p('h(u) = mean of item embeddings over H(u)\nr(u) = γ × user_embedding(u) + (1 − γ) × W × h(u)')
p('The mixing weight γ is fixed at 0.5 and W is a learned projection without bias. Empty histories have a zero mean. History and target items share the embedding table, allowing gradients to update history embeddings. Dropout applies to the combined representation; normalized user and item vectors produce cosine scores.')
p('For positive score s+ and negative score s−, the per-pair objective is 1 − s+ + w × max(0, s− − m), where m is the margin and w the negative weight. The implementation averages this over sampled pairs. SimpleX regularization is zero in the selected recipes. The BPR command-line selector supplies pairwise samples; SimpleX computes CCL internally.')
p('Training-only histories are required at construction and copied into compact embedding bags. During tuning, validation interactions are excluded. After parameter selection, the former validation portion joins final training; test interactions remain excluded. The shared fitting helper also creates LightGCN’s interaction graph before model construction. LightGCN retains DaisyRec’s normalized bipartite adjacency, sparse layer propagation and equal averaging of layer-zero and propagated embeddings. Public SciPy sparse construction replaces an obsolete private API. LightGCN is regression-tested but is not one of the three final comparison models.')
h('5 Hyperparameter tuning')
p('Validation NDCG@10 selects the best complete configuration. Each model completed four trials on MovieLens and eight on Yelp, for 36 trials in total. The Optuna TPE sampler uses ten startup trials, so these short studies remain in its random startup phase; they are not completed adaptive Bayesian searches. All runs use seed 2022, four uniformly sampled negatives per positive and no early stopping. MF runs for 20 epochs; NeuMF and SimpleX run for 30.')
p('Table 2. Search spaces used for both datasets unless indicated', 'Caption')
table(['Parameter','MF','NeuMF','SimpleX'],[
    ['Factors','32, 64, 96, 128','16, 32, 48, 64','32, 64, 96, 128'],
    ['Learning rate','0.001–0.046; step 0.005','0.0005–0.010; step 0.0005','0.0005–0.010; step 0.0005'],
    ['L1 and L2','0–0.010; step 0.001','0–0.010; step 0.001','Fixed at 0'],
    ['Batch size','256, 512, 1024','256, 512','ML: 256, 512; Yelp: 128, 256, 512'],
    ['Other','—','Layers: 1, 2, 3; dropout: 0–0.5, step 0.1','Margin: 0.5, 0.7, 0.9; weight: 5, 10, 20; dropout: 0–0.3, step 0.1']],
    [1.05,1.8,1.85,2.05],9)
p('MF’s configured upper learning-rate bound is 0.05; the stepped sampler’s last reachable value is 0.046. Defaults resolve to SGD with normal initialization for MF, Adam with Xavier normal for NeuMF, and Adam with Xavier uniform for SimpleX.')

page()
h('6 Tuning evidence and selected settings')
p('Figure 1 shows validation NDCG@10 for every completed trial on both datasets. The outlined point identifies the selected trial. The plots label trials from one, while saved trial IDs start at zero. Per-trial scores, durations and parameters are retained in trials.csv, complete recipes in trial_configs/, and resumable studies in study.db.')
pic = p().add_run().add_picture(str(ROOT/'artifacts/verified/plots/ndcg_vs_trial.png'),width=Inches(6.55))
pic._inline.docPr.set('descr','Six panels of per-trial validation NDCG at 10 for MF, NeuMF and SimpleX on MovieLens and Yelp, showing four and eight completed trials respectively.')
p('Figure 1. Completed tuning trials; source: artifacts/verified/*/trials.csv.', 'Caption')
p('Table 3. Selected settings and validation NDCG@10', 'Caption')
with (ROOT/'artifacts/verified/selected_parameters.csv').open() as f:
    selected=list(csv.DictReader(f))
labels={'mf':'MF','neumf':'NeuMF','simplex':'SimpleX'}
table(['Model / dataset','Trial ID','Factors','Batch','LR','L1 / L2','Val. NDCG'],[
    [labels[r['model']]+' / '+r['dataset'],r['selected_trial'],r['factors'],r['batch_size'],f"{float(r['lr']):g}",f"{float(r['reg_1']):g} / {float(r['reg_2']):g}",f"{float(r['validation_ndcg']):.6f}"] for r in selected],
    [1.8,.55,.55,.55,.65,.95,1.05],8.5)
p('NeuMF selects 3 layers / 0.2 dropout for MovieLens and 2 / 0.3 for Yelp. SimpleX selects margin / negative weight / dropout of 0.9 / 20 / 0.2 for MovieLens and 0.5 / 10 / 0.3 for Yelp. Both use γ = 0.5. Full-precision settings, including fixed values, are loaded from best_config.json in every final test.')

page()
h('7 Final evaluation and discussion')
p('Ranking uses 1,000 sampled candidates per user, includes ground-truth items, and excludes training items from candidate negatives. DaisyRec’s candidate sampler retains sampling with replacement. Metrics are evaluated at K = 10; these are sampled-ranking results, not full-catalog estimates. The global split includes 660 MovieLens and 1,155 Yelp test users without training history.')
with (ROOT/'artifacts/verified/final_summary.csv').open() as f:
    results=list(csv.DictReader(f))
for dataset, caption in [('ml-50k','Table 4. MovieLens final ranking results'),('yelp-200k','Table 5. Yelp final ranking results')]:
    p(caption,'Caption')
    table(['Model','Recall','NDCG','MRR','Hit ratio','Precision'],[
        [labels[r['model']]]+[f"{float(r[k]):.6f}" for k in ['Recall','NDCG','MRR','Hit Ratio','Precision']]
        for r in results if r['dataset']==dataset],[1.1,1.1,1.1,1.1,1.1,1.25],9.5)
p('NeuMF leads all five ranking metrics on MovieLens, with NDCG@10 of 0.120102 versus 0.086964 for MF and 0.040691 for SimpleX. On Yelp, SimpleX leads all five metrics, reaching NDCG@10 of 0.082864 versus 0.032882 for MF and 0.018919 for NeuMF. Thus no model dominates across both datasets under the same evaluation protocol. Architecture, selected settings and dataset characteristics may all contribute; these experiments do not isolate their causal effects.')
p('Table 6. Recommendation exposure at K = 10', 'Caption')
table(['Model / dataset','Coverage','Gini','Entropy in bits'],[
    [labels[r['model']]+' / '+r['dataset']]+[f"{float(r[k]):.6f}" for k in ['Coverage','Gini','Entropy']] for r in results],
    [2.4,1.3,1.3,1.75],9.5)
p('Coverage measures the fraction of catalog items recommended. Entropy measures the distribution of recommendation exposure. The retained Gini calculation includes only items receiving at least one recommendation, so it is not a catalog-wide Gini. On Yelp, SimpleX combines the strongest accuracy with the highest coverage (0.925206), highest entropy and lowest Gini among the compared models. On MovieLens, NeuMF’s stronger ranking accuracy accompanies lower coverage (0.330213) and more concentrated exposure. These measures do not establish user satisfaction, fairness or within-list diversity.')
p('All six final metric CSVs and manifests are saved in artifacts/verified/results/<dataset>/. Each test uses its selected training recipe; only the reporting metric list is extended. The automated audit verifies complete trials, selected settings, package versions, Linux provenance, current code/data hashes and finite metric values.')

page()
h('8 Reproducibility and verification')
p('Experiments ran on Ubuntu 24.04 under WSL2 with Python 3.12.3, CPU PyTorch 2.5.1, pandas 2.2.3, NumPy 1.26.4, SciPy 1.14.1 and Optuna 4.1.0. Pandas and PyTorch imports succeeded in Linux. Five behavioral regression tests passed, covering history aggregation and gradients, held-out-item isolation, copied and empty histories, graph-dependent LightGCN propagation, fitting-split construction, single-user ranking and loader equivalence. Linux checks and the six-result audit passed again on 6 October 2026.')
p('The README documents setup, dataset preparation and individual tuning/test commands. From the project root in the activated Linux environment:')
code('python scripts/run_experiments.py --stage all --jobs 2\npython scripts/audit_results.py')
p('Completed studies resume without discarding trials. To repeat only final evaluation with the saved selected configurations:')
code('python scripts/run_experiments.py --stage test --jobs 2')
p('Final evaluation fits each model again before testing. The workflow saves configurations, metrics and provenance, not trained model checkpoints. Evidence includes final_summary.csv, selected_parameters.csv, tuning_trials.csv, audit.json, per-run logs, per-trial recipes and both per-trial and cumulative-best tuning plots under artifacts/verified/. Historical and interrupted outputs are stored separately.')
h('9 Limitations and conclusion')
p('The short trial budgets and single seed limit confidence in model ordering. The studies do not provide statistical significance, repeated-seed uncertainty, a training-history ablation, or exhaustive optimization. The sparse, partly cold-start data and candidate sampling with replacement further constrain generalization. Four negatives per positive also differ from the large-negative-budget SimpleX benchmarks [4]. The historical Yelp release is uncertain; the supplied subset and its hashes are necessary for exact reproduction.')
p('The project integrates training-history-aware SimpleX, repairs split-dependent graph construction for LightGCN, completes all configured tuning studies and evaluates six selected recipes. NeuMF performs best on the MovieLens subset and SimpleX on the Yelp subset in these runs. Future work should increase tuning and negative-sampling budgets, repeat experiments across seeds, compare warm-user and full-catalog protocols, and ablate history aggregation before attributing performance gains to a specific component.')
h('References')
refs=[
    '[1] DaisyRec contributors. DaisyRec v2.0 repository and framework. https://github.com/recsys-benchmark/DaisyRec-v2.0',
    '[2] GroupLens Research. MovieLens 1M Dataset. https://grouplens.org/datasets/movielens/1m/',
    '[3] Yelp. Yelp Dataset on Kaggle. https://www.kaggle.com/datasets/yelp-dataset/yelp-dataset/data',
    '[4] Mao, K., Zhu, J., Wang, J., Dai, Q., Dong, Z., Xiao, X. and He, X. (2021). SimpleX: A Simple and Strong Baseline for Collaborative Filtering. CIKM. https://arxiv.org/abs/2109.12613',
]
for ref in refs:
    para=p(ref)
    para.paragraph_format.space_after=Pt(5)
    for run in para.runs:
        run.font.size=Pt(9)
doc.save(OUT)
print(OUT)
