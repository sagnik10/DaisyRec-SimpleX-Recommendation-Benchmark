"""Convert personal presentation notes to a readable Word document."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.65)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.bold = False
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',22),('Heading 1',15)]:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

def inline(p, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if token.startswith('**'):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith('`'):
            r = p.add_run(token[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            p.add_run(token)

def table(data):
    # Keep tuning numbers and model-specific settings in two legible tables.
    if len(data[0]) == 8:
        table([row[:7] for row in data])
        table([[data[0][0], data[0][-1]]] + [[row[0], row[-1]] for row in data[1:]])
        return
    t = doc.add_table(rows=0, cols=len(data[0]))
    t.autofit = False
    widths = {7:[1.0,.85,.9,.65,.65,1.0,1.0], 5:[1.0,1.1,1.2,1.25,2.0], 2:[1.65,5.3]}[len(data[0])]
    # Scale all widths to the available page width.
    widths = [w * 6.97 / sum(widths) for w in widths]
    for ri, values in enumerate(data):
        row = t.add_row()
        pr = row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        if ri == 0: pr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[ci])
            cell.text = value
            tcpr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E9EEF3' if ri == 0 else ('F7F8FA' if ri % 2 == 0 else 'FFFFFF')); tcpr.append(sh)
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b = OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            tcpr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
                for r in p.runs: r.font.size = Pt(9); r.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = (ROOT/'docs/PRESENTATION_NOTES.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if line.startswith('|'):
        data = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', c) for c in cells): data.append(cells)
            i += 1
        table(data); continue
    if line.startswith('# '):
        doc.add_paragraph('DaisyRec Presentation Notes', 'Title')
    elif line.startswith('## '):
        doc.add_heading(line[3:],1)
    elif line:
        p = doc.add_paragraph()
        if line.startswith('- '):
            p.style = 'List Bullet'; line = line[2:]
        elif re.match(r'\d+\. ',line):
            p.paragraph_format.left_indent = Inches(.2)
            p.paragraph_format.first_line_indent = Inches(-.2)
        inline(p,line)
    i += 1
doc.core_properties.title = 'DaisyRec Presentation Notes'
doc.core_properties.subject = 'Personal presentation preparation'
out = ROOT/'DaisyRec Presentation Notes.docx'
doc.save(out)
print(out.name)
